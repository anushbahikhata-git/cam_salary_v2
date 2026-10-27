import sys
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.workbook import WorkbookError


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float):
        return str(round(value, 4))
    return str(value).replace("\n", " ").strip()


def _row_text(row: tuple[Any, ...]) -> str | None:
    values = [_cell(value) for value in row]
    if not any(values):
        return None
    return " | ".join(values)


def read_msme_workbook(data: bytes):
    try:
        return load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise WorkbookError("Upload a valid .xlsx workbook.") from exc


def validate_workbook(data: bytes) -> None:
    read_msme_workbook(data).close()


def workbook_to_text(workbook) -> str:
    required_sheets = [
        "Summary",
        "Overview",
        "Transactions",
    ]

    excluded_transaction_columns = {
        "Particulars",
        "Bank name",
        "Account type",
        "Account number",
        "Cheque/Reference number",
    }

    parts = []

    try:
        for sheet_name in required_sheets:
            if sheet_name not in workbook.sheetnames:
                continue

            sheet = workbook[sheet_name]
            rows = iter(sheet.iter_rows(values_only=True))

            if sheet_name == "Transactions":
                headers = next(rows, None)

                if headers is None:
                    continue

                excluded_indexes = {
                    index
                    for index, header in enumerate(headers)
                    if header is not None
                    and _cell(header) in excluded_transaction_columns
                }

                kept_indexes = [
                    index
                    for index in range(len(headers))
                    if index not in excluded_indexes
                ]

                header_values = [_cell(headers[index]) for index in kept_indexes]

                if any(header_values):
                    parts.append(" | ".join(header_values) + "\n")

                for row in rows:
                    values = [
                        _cell(row[index]) if index < len(row) else ""
                        for index in kept_indexes
                    ]

                    if any(values):
                        parts.append(" | ".join(values) + "\n")

            else:
                for row in rows:
                    line = _row_text(row)
                    if line is not None:
                        parts.append(line + "\n")

    finally:
        workbook.close()

    return "".join(parts)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit("usage: python -m app.msme_workbook <report.xlsx> [output.txt]")

    with open(sys.argv[1], "rb") as handle:
        extracted_text = workbook_to_text(read_msme_workbook(handle.read()))

    output_path = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("tests") / "msme.txt"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(extracted_text, encoding="utf-8")

    print("Total characters:", len(extracted_text))
    print("Saved to:", output_path.resolve())
