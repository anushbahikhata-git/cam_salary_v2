from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.config import MAX_TEXT_CHARS


class WorkbookError(Exception):
    pass


class WorkbookTooLargeError(WorkbookError):
    pass


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float):
        return str(round(value, 4))
    return str(value).replace("\n", " ").strip()


def _priority(title: str) -> tuple[int, str]:
    name = title.casefold()
    if any(term in name for term in ("overview", "summary", "average balance")):
        return 0, name
    if any(term in name for term in ("transaction", "statement", "bank", "credit", "debit")):
        return 1, name
    return 2, name


def _row_text(row: tuple[Any, ...]) -> str | None:
    values = [_cell(value) for value in row]
    if not any(values):
        return None
    return " | ".join(values)


def to_text(data: bytes) -> str:
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise WorkbookError("Upload a valid .xlsx workbook.") from exc

    parts = []
    used = 0

    def append(line: str) -> None:
        nonlocal used
        if used + len(line) > MAX_TEXT_CHARS:
            raise WorkbookTooLargeError(
                "Workbook content exceeds the text processing limit. "
                "Reduce or split the workbook and retry."
            )
        parts.append(line)
        used += len(line)

    try:
        for sheet in sorted(book.worksheets, key=lambda item: _priority(item.title)):
            rows = iter(sheet.iter_rows(values_only=True))
            first_row = None
            for row in rows:
                first_row = _row_text(row)
                if first_row is not None:
                    break
            if first_row is None:
                continue

            append(f"### SHEET: {sheet.title} | rows={sheet.max_row} | cols={sheet.max_column}\n")
            append(f"{first_row}\n")
            for row in rows:
                line = _row_text(row)
                if line is not None:
                    append(f"{line}\n")
    except WorkbookError:
        raise
    except Exception as exc:
        raise WorkbookError("The workbook could not be read.") from exc
    finally:
        book.close()

    return "".join(parts)
